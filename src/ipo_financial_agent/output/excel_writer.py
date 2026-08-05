from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ipo_financial_agent.models import (
    FinancialNote,
    MetricResult,
    RawStatementTable,
    RiskFinding,
    StatementFact,
)

DARK_BLUE = "17365D"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
RED = "FF0000"
BLUE = "0000FF"
GREEN = "008000"
THIN = Side(style="thin", color="B7B7B7")
NUMBER_FORMAT = '#,##0.00;[Red](#,##0.00);-'
PERCENT_FORMAT = '0.0%;[Red](0.0%);-'


def _safe_sheet_name(name: str, used: set[str]) -> str:
    clean = re.sub(r"[\\/*?:\[\]]", "_", name).strip() or "未命名表"
    clean = clean[:31]
    candidate = clean
    index = 2
    while candidate in used:
        suffix = f"_{index}"
        candidate = clean[: 31 - len(suffix)] + suffix
        index += 1
    used.add(candidate)
    return candidate


def _apply_header(row) -> None:
    for cell in row:
        cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)


def _set_widths(sheet, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _write_raw_statement_sheet(sheet, table: RawStatementTable) -> None:
    max_cols = max((len(row) for row in table.rows), default=1)
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_cols + 1)
    title = sheet.cell(1, 1, table.statement_name)
    title.font = Font(size=14, bold=True)
    title.alignment = Alignment(horizontal="center")
    sheet.cell(2, 1, f"公司：{table.company}")
    sheet.cell(2, 2, f"单位：{table.unit or '未识别'}")
    sheet.cell(2, 3, f"币种：{table.currency or '未识别'}")
    sheet.cell(2, 4, f"口径：{table.entity_scope or '未识别'}")
    sheet.cell(3, 1, f"来源页码：{', '.join('P'+str(page) for page in table.pages)}")

    start_row = 5
    for row_index, raw_row in enumerate(table.rows, start=start_row):
        padded = raw_row + [""] * (max_cols - len(raw_row))
        page = table.row_pages[row_index - start_row] if row_index - start_row < len(table.row_pages) else None
        for column_index, value in enumerate(padded, start=1):
            cell = sheet.cell(row_index, column_index, value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=THIN)
            if page is not None:
                cell.comment = Comment(f"来源页码：P{page}", "IPO Financial Agent")
        sheet.cell(row_index, max_cols + 1, f"P{page}" if page is not None else "")

    # 第一行通常是表头，做轻量样式，不强行假设所有表结构一致。
    if table.rows:
        for cell in sheet[start_row]:
            cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheet.cell(start_row - 1, max_cols + 1, "来源页码")
    sheet.cell(start_row - 1, max_cols + 1).fill = PatternFill("solid", fgColor=DARK_BLUE)
    sheet.cell(start_row - 1, max_cols + 1).font = Font(color=WHITE, bold=True)
    sheet.freeze_panes = f"A{start_row}"
    sheet.column_dimensions["A"].width = 34
    for col in range(2, max_cols + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 17
    sheet.column_dimensions[get_column_letter(max_cols + 1)].width = 14


def export_financial_workbook(
    *,
    output_path: str | Path,
    raw_statements: list[RawStatementTable],
    facts: list[StatementFact],
    notes: list[FinancialNote],
    metrics: list[MetricResult],
    risks: list[RiskFinding],
) -> Path:
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()

    directory = workbook.create_sheet(_safe_sheet_name("目录", used))
    directory.append(["序号", "类型", "名称", "页码", "Sheet"])
    _apply_header(directory[1])
    directory_index = 1

    for table in raw_statements:
        sheet_name = _safe_sheet_name(table.statement_name, used)
        sheet = workbook.create_sheet(sheet_name)
        _write_raw_statement_sheet(sheet, table)
        directory_index += 1
        directory.append([
            directory_index - 1,
            "原样财务报表",
            table.statement_name,
            ", ".join(f"P{page}" for page in table.pages),
            sheet_name,
        ])

    fact_sheet = workbook.create_sheet(_safe_sheet_name("财务事实明细", used))
    fact_headers = [
        "fact_id", "报表名称", "原始科目", "计算标签", "期间", "金额", "原始值",
        "单位", "币种", "页码", "附注", "主体口径", "来源表", "置信度",
    ]
    fact_sheet.append(fact_headers)
    _apply_header(fact_sheet[1])
    for fact in facts:
        fact_sheet.append([
            fact.fact_id, fact.statement_name, fact.item_name, fact.canonical_tag,
            fact.period, fact.value, fact.raw_value, fact.unit, fact.currency,
            fact.page, fact.note, fact.entity_scope, fact.source_table_id, fact.confidence,
        ])
        amount_cell = fact_sheet.cell(fact_sheet.max_row, 6)
        amount_cell.number_format = NUMBER_FORMAT
        amount_cell.font = Font(color=BLUE)
        amount_cell.comment = Comment(f"来源页码：P{fact.page}", "IPO Financial Agent")
    fact_sheet.freeze_panes = "A2"
    fact_sheet.sheet_view.showGridLines = False
    _set_widths(fact_sheet, {"A": 24, "B": 28, "C": 34, "D": 22, "E": 18, "F": 16, "G": 16, "H": 16, "I": 12, "J": 10, "K": 14, "L": 18, "M": 24, "N": 12})

    note_sheet = workbook.create_sheet(_safe_sheet_name("重点财务资料", used))
    note_headers = ["note_id", "主题", "标题", "信息类型", "摘要", "页码", "原文摘录", "置信度", "表格数量"]
    note_sheet.append(note_headers)
    _apply_header(note_sheet[1])
    for note in notes:
        note_sheet.append([
            note.note_id, note.topic, note.title, note.information_type, note.summary,
            ", ".join(f"P{page}" for page in note.pages), note.source_excerpt,
            note.confidence, len(note.tables),
        ])
    note_sheet.freeze_panes = "A2"
    note_sheet.sheet_view.showGridLines = False
    _set_widths(note_sheet, {"A": 24, "B": 18, "C": 30, "D": 20, "E": 54, "F": 18, "G": 54, "H": 12, "I": 12})
    for row in note_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    metric_sheet = workbook.create_sheet(_safe_sheet_name("财务指标", used))
    metric_sheet.append(["指标", "代码", "期间", "结果", "数值", "公式", "来源页码", "状态"])
    _apply_header(metric_sheet[1])
    for metric in metrics:
        metric_sheet.append([
            metric.metric_name, metric.metric_code, metric.period, metric.display_value,
            metric.value, metric.formula, ", ".join(f"P{page}" for page in metric.source_pages), metric.status,
        ])
        metric_sheet.cell(metric_sheet.max_row, 5).font = Font(color=GREEN)
        if "率" in metric.metric_name or "margin" in metric.metric_code or "growth" in metric.metric_code or metric.metric_code == "debt_ratio":
            metric_sheet.cell(metric_sheet.max_row, 5).number_format = PERCENT_FORMAT
        else:
            metric_sheet.cell(metric_sheet.max_row, 5).number_format = "0.00x"
    metric_sheet.freeze_panes = "A2"
    metric_sheet.sheet_view.showGridLines = False
    _set_widths(metric_sheet, {"A": 24, "B": 26, "C": 18, "D": 14, "E": 14, "F": 24, "G": 18, "H": 16})

    risk_sheet = workbook.create_sheet(_safe_sheet_name("风险事项", used))
    risk_sheet.append(["类别", "标题", "严重程度", "说明", "来源页码", "规则代码"])
    _apply_header(risk_sheet[1])
    for risk in risks:
        risk_sheet.append([
            risk.category, risk.title, risk.severity, risk.description,
            ", ".join(f"P{page}" for page in risk.source_pages), risk.rule_code,
        ])
        if risk.severity == "high":
            risk_sheet.cell(risk_sheet.max_row, 3).fill = PatternFill("solid", fgColor="F4CCCC")
        elif risk.severity == "medium":
            risk_sheet.cell(risk_sheet.max_row, 3).fill = PatternFill("solid", fgColor="FFF2CC")
    risk_sheet.freeze_panes = "A2"
    risk_sheet.sheet_view.showGridLines = False
    _set_widths(risk_sheet, {"A": 18, "B": 32, "C": 14, "D": 64, "E": 18, "F": 24})
    for row in risk_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    evidence_sheet = workbook.create_sheet(_safe_sheet_name("证据索引", used))
    evidence_sheet.append(["证据类型", "ID", "主题/科目", "页码", "内容"])
    _apply_header(evidence_sheet[1])
    for fact in facts:
        evidence_sheet.append(["财务事实", fact.fact_id, fact.item_name, f"P{fact.page}", f"{fact.period}: {fact.raw_value} {fact.unit or ''}".strip()])
    for note in notes:
        evidence_sheet.append(["财务附注", note.note_id, note.topic, ", ".join(f"P{page}" for page in note.pages), note.summary])
    evidence_sheet.freeze_panes = "A2"
    evidence_sheet.sheet_view.showGridLines = False
    _set_widths(evidence_sheet, {"A": 16, "B": 24, "C": 30, "D": 18, "E": 70})
    for row in evidence_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    _set_widths(directory, {"A": 10, "B": 20, "C": 36, "D": 20, "E": 30})
    directory.freeze_panes = "A2"
    directory.sheet_view.showGridLines = False

    workbook.save(target)
    return target
