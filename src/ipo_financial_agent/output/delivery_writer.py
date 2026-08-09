"""Build the stable, user-facing IPO due diligence delivery bundle."""

from __future__ import annotations

from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill


_HEADER_FILL = PatternFill("solid", fgColor="17365D")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_NON_XML_RE = re.compile(
    r"[^\x09\x0A\x0D\x20-\uD7FF\uE000-\uFFFD\U00010000-\U0010FFFF]"
)


def _excel_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    # Web extraction can return PDF/binary control characters. Excel rejects
    # them and limits a cell to 32,767 characters.
    cleaned = ILLEGAL_CHARACTERS_RE.sub("", value)
    return _NON_XML_RE.sub("", cleaned)[:32767]


def _append(sheet: Any, values: list[Any]) -> None:
    sheet.append([_excel_value(value) for value in values])


def _value(item: Any, name: str, default: Any = None) -> Any:
    if isinstance(item, dict):
        return item.get(name, default)
    return getattr(item, name, default)


def _prepare_sheet(workbook: Any, title: str, headers: list[str]) -> Any:
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    _append(sheet, headers)
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    return sheet


def _finish_sheet(sheet: Any, widths: list[float]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def enrich_due_diligence_workbook(
    *,
    workbook_path: str | Path,
    evidence: list[Any],
    findings: list[Any],
    diligence_questions: list[Any],
    agent_messages: list[Any],
    conclusion: Any,
    report_review: Any,
) -> Path:
    """Append the non-financial Agent artifacts required by the product PRD."""
    target = Path(workbook_path)
    workbook = load_workbook(target)

    summary = _prepare_sheet(workbook, "尽调结论", ["项目", "结果"])
    for label, value in (
        ("尽调结论", _value(conclusion, "verdict", "")),
        ("历史财务质量", _value(conclusion, "historical_financial_quality", "")),
        ("未来盈利能力", _value(conclusion, "future_earning_power", "")),
        ("重大风险等级", _value(conclusion, "material_risk_level", "")),
        ("结论置信度", _value(conclusion, "confidence", "")),
        ("报告终审是否通过", _value(report_review, "passed", False)),
        ("报告终审得分", _value(report_review, "score", 0)),
        ("终审摘要", _value(report_review, "summary", "")),
    ):
        _append(summary, [label, value])
    _finish_sheet(summary, [24, 90])

    evidence_sheet = _prepare_sheet(
        workbook,
        "Research Evidence",
        ["Evidence ID", "来源类型", "标题", "页码", "URL", "内容", "置信度"],
    )
    for item in evidence:
        _append(
            evidence_sheet,
            [
                _value(item, "evidence_id", ""),
                _value(item, "source_type", ""),
                _value(item, "title", ""),
                _value(item, "page_number", ""),
                _value(item, "source_url", ""),
                _value(item, "content", ""),
                _value(item, "confidence", ""),
            ]
        )
    _finish_sheet(evidence_sheet, [24, 16, 32, 10, 48, 90, 12])

    finding_sheet = _prepare_sheet(
        workbook,
        "Research Findings",
        ["Finding ID", "Agent", "研究问题", "结论", "Evidence IDs", "证据强度", "风险", "待核事项"],
    )
    for item in findings:
        _append(
            finding_sheet,
            [
                _value(item, "finding_id", ""),
                _value(item, "agent_name", ""),
                _value(item, "question", ""),
                _value(item, "conclusion", ""),
                ", ".join(_value(item, "evidence_ids", []) or []),
                _value(item, "evidence_strength", ""),
                "；".join(_value(item, "risks", []) or []),
                "；".join(_value(item, "open_questions", []) or []),
            ]
        )
    _finish_sheet(finding_sheet, [24, 22, 46, 90, 52, 14, 42, 56])

    question_sheet = _prepare_sheet(
        workbook,
        "补充尽调清单",
        ["优先级", "类别", "问题", "原因", "所需材料", "未解决影响", "状态"],
    )
    for item in diligence_questions:
        _append(
            question_sheet,
            [
                _value(item, "priority", ""),
                _value(item, "category", ""),
                _value(item, "question", ""),
                _value(item, "rationale", ""),
                "；".join(_value(item, "requested_materials", []) or []),
                _value(item, "downside_if_unresolved", ""),
                _value(item, "status", ""),
            ]
        )
    _finish_sheet(question_sheet, [10, 22, 60, 60, 56, 56, 12])

    trace_sheet = _prepare_sheet(
        workbook,
        "Agent Trace",
        ["时间", "发送方", "接收方", "类型", "内容", "结构化载荷"],
    )
    for item in agent_messages:
        _append(
            trace_sheet,
            [
                _value(item, "timestamp", ""),
                _value(item, "sender", ""),
                _value(item, "receiver", ""),
                _value(item, "message_type", ""),
                _value(item, "content", ""),
                str(_value(item, "payload", {}) or {}),
            ]
        )
    _finish_sheet(trace_sheet, [24, 24, 24, 14, 90, 70])

    # Final boundary guard also covers values loaded from the financial
    # workbook and future delivery sheets added outside `_append`.
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = _excel_value(cell.value)
    workbook.save(target)
    return target


def ledger_integrity(evidence: list[Any], findings: list[Any]) -> dict[str, Any]:
    """Return an auditable Evidence -> Finding integrity result."""
    evidence_ids = [_value(item, "evidence_id", "") for item in evidence]
    known = {item for item in evidence_ids if item}
    duplicates = sorted({item for item in evidence_ids if evidence_ids.count(item) > 1})
    missing: dict[str, list[str]] = {}
    for finding in findings:
        absent = [
            item
            for item in (_value(finding, "evidence_ids", []) or [])
            if item not in known
        ]
        if absent:
            missing[_value(finding, "finding_id", "unknown")] = absent
    return {
        "passed": not duplicates and not missing,
        "evidence_count": len(evidence_ids),
        "finding_count": len(findings),
        "duplicate_evidence_ids": duplicates,
        "missing_evidence_references": missing,
    }
