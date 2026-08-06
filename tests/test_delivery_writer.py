from pathlib import Path

from openpyxl import Workbook, load_workbook

from ipo_financial_agent.models_agent import (
    AgentMessage,
    DiligenceQuestion,
    DueDiligenceConclusion,
    Evidence,
    Finding,
    ReportReview,
)
from ipo_financial_agent.output.delivery_writer import (
    enrich_due_diligence_workbook,
    ledger_integrity,
)


def _evidence_and_finding() -> tuple[Evidence, Finding]:
    evidence = Evidence(
        source_type="prospectus",
        page_number=10,
        title="收入披露",
        content="报告期收入已披露。",
    )
    finding = Finding(
        agent_name="financial_dd",
        question="收入是否已披露？",
        conclusion="已披露。",
        evidence_ids=[evidence.evidence_id],
    )
    return evidence, finding


def test_ledger_integrity_checks_finding_references() -> None:
    evidence, finding = _evidence_and_finding()
    assert ledger_integrity([evidence], [finding])["passed"] is True

    broken = finding.model_copy(update={"evidence_ids": ["ev_missing"]})
    result = ledger_integrity([evidence], [broken])
    assert result["passed"] is False
    assert result["missing_evidence_references"][finding.finding_id] == ["ev_missing"]


def test_due_diligence_workbook_contains_product_closure_sheets(tmp_path: Path) -> None:
    evidence, finding = _evidence_and_finding()
    target = tmp_path / "IPO_Due_Diligence_Report.xlsx"
    workbook = Workbook()
    workbook.active.title = "财务底稿"
    workbook.save(target)

    enrich_due_diligence_workbook(
        workbook_path=target,
        evidence=[evidence],
        findings=[finding],
        diligence_questions=[
            DiligenceQuestion(
                priority="P0",
                category="financial",
                question="核验收入真实性。",
                rationale="现金流匹配不足。",
            )
        ],
        agent_messages=[
            AgentMessage(
                sender="ResearchManager",
                receiver="FinancialAgent",
                message_type="plan",
                content="核验收入。",
            )
        ],
        conclusion=DueDiligenceConclusion(verdict="conditional_proceed"),
        report_review=ReportReview(passed=True, score=100),
    )

    rendered = load_workbook(target)
    for sheet in (
        "财务底稿",
        "尽调结论",
        "Research Evidence",
        "Research Findings",
        "补充尽调清单",
        "Agent Trace",
    ):
        assert sheet in rendered.sheetnames
    assert rendered["Research Evidence"]["A2"].value == evidence.evidence_id
