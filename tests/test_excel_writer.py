from pathlib import Path

from openpyxl import load_workbook

from ipo_financial_agent.models import RawStatementTable
from ipo_financial_agent.output.excel_writer import export_financial_workbook


def test_excel_contains_raw_statement_and_page(tmp_path: Path):
    table = RawStatementTable(
        table_id="t1",
        statement_name="综合财务状况表",
        statement_type="balance_sheet",
        company="测试公司",
        pages=[10],
        rows=[["项目", "2024"], ["现金", "100"]],
        row_pages=[10, 10],
        source_file="test.pdf",
    )
    output = tmp_path / "out.xlsx"
    export_financial_workbook(
        output_path=output,
        raw_statements=[table],
        facts=[],
        notes=[],
        metrics=[],
        risks=[],
    )
    workbook = load_workbook(output)
    assert "综合财务状况表" in workbook.sheetnames
    sheet = workbook["综合财务状况表"]
    assert sheet["A6"].value == "现金"
    assert sheet["C6"].value == "P10"
